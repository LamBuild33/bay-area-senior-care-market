#!/usr/bin/env python3
"""Turn a week of visit notes into a weekly invoice.

You paste the visit notes (the ones written with docs/visit_note_form.md) into a
text file. This tool reads them, does the arithmetic with exact decimal math, shows
you what it understood, and, only after you say yes, writes a plain-text invoice you
can paste into an email. It can also add a row to the invoice tracker spreadsheet.

The money rules come from the client agreement (sections 5 to 8) and the README:
  * Bill the actual time, arrival to departure, to the nearest minute.
  * The shortest visit for that client still applies.
  * Mileage is for drives with or for the client only, at the mileage rate.
  * Purchases on the client's card are listed but not added.
  * Purchases on my own card are added (backup only, under $100 each).
  * Payment is due 7 days after the invoice is sent.

Nothing here is legal or tax advice. The tool does arithmetic; you check it.
Private files (clients.json, notes, invoices) must stay out of the GitHub repo.

Usage:
  python visit_to_invoice.py notes.txt --client example --clients clients.json
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

CENT = Decimal("0.01")
MAX_DAY_MINUTES = 12 * 60          # agreement: up to 12 hours in a day
OWN_CARD_LIMIT = Decimal("100")    # agreement section 7: own card only under $100
DEFAULT_MILEAGE_RATE = Decimal("0.76")  # IRS standard rate, July to December 2026


def money(x):
    return Decimal(x).quantize(CENT, ROUND_HALF_UP)


def fmt(x):
    return f"${money(x):,.2f}"


def hm(minutes):
    h, m = divmod(int(minutes), 60)
    return f"{h} h {m:02d} min"


def clock(minutes):
    h, m = divmod(int(minutes), 60)
    suffix = "am" if h < 12 else "pm"
    return f"{(h % 12) or 12}:{m:02d} {suffix}"


def pretty_date(d):
    return f"{d.strftime('%b')} {d.day}, {d.year}"


# --------------------------------------------------------------------------
# Reading the notes
# --------------------------------------------------------------------------

LABELS = [
    ("date", r"date"),
    ("times", r"start time\s*[-–—]\s*end time|start\s*/\s*end time|times?"),
    ("did", r"what we did"),
    ("managed", r"what you managed on your own"),
    ("noticed", r"anything i noticed"),
    ("drives", r"drives or purchases[^:]*"),
    ("next", r"next visit"),
]
LABEL_RES = [(k, re.compile(rf"^\s*(?:{p})\s*:\s*(.*)$", re.I)) for k, p in LABELS]
NOTE_START = re.compile(r"^\s*visit note\b", re.I)
TIME_RE = re.compile(r"(\d{1,2})(?::(\d{2}))?\s*([ap])\.?\s*m?\.?", re.I)
MILES_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:miles?|mi)\b", re.I)
DOLLAR_RE = re.compile(r"\$\s*(\d+(?:,\d{3})*(?:\.\d{1,2})?)")
SENTENCE_SPLIT = re.compile(r"\.\s+(?=[A-Z])|\n|;\s*")
CLIENT_CARD = ("your card", "client's card", "client card", "their card",
               "your debit", "your credit")
MY_CARD = ("my card", "my own card", "own card", "my debit", "my credit")


@dataclass
class Purchase:
    date: date
    text: str
    amount: Decimal
    card: str  # "client", "mine" or "unknown"


@dataclass
class Visit:
    day: date = None
    start: int = None          # minutes after midnight
    end: int = None
    actual: int = None         # minutes worked
    miles: Decimal = Decimal("0")
    purchases: list = field(default_factory=list)
    errors: list = field(default_factory=list)    # block writing the invoice
    warnings: list = field(default_factory=list)  # shown, do not block
    label: str = ""            # how to point at this note in messages


def parse_date(raw, year):
    s = raw.strip()
    weekday = None
    m = re.match(r"^([A-Za-z]+day),?\s*(.*)$", s)
    if m:
        weekday, s = m.group(1), m.group(2)
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s, flags=re.I)
    s = s.replace(",", " ").replace("Sept ", "Sep ")
    s = re.sub(r"\s+", " ", s).strip()
    tries = []
    if re.search(r"\b\d{4}\b", s) or re.search(r"\d+/\d+/\d+", s):
        tries += [(s, "%B %d %Y"), (s, "%b %d %Y"), (s, "%m/%d/%Y")]
    else:
        tries += [(f"{s} {year}", "%B %d %Y"), (f"{s} {year}", "%b %d %Y"),
                  (f"{s}/{year}", "%m/%d/%Y")]
    for text, f in tries:
        try:
            return datetime.strptime(text, f).date(), weekday
        except ValueError:
            continue
    return None, weekday


def parse_times(raw):
    found = []
    for m in TIME_RE.finditer(raw):
        h, mi, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3).lower()
        if not (1 <= h <= 12 and 0 <= mi < 60):
            return None
        found.append((h % 12) * 60 + mi + (720 if ap == "p" else 0))
    return found if len(found) == 2 else None


def parse_drives(text, day):
    miles_found = [Decimal(x) for x in MILES_RE.findall(text)]
    purchases, notes = [], []
    for sentence in SENTENCE_SPLIT.split(text):
        s = sentence.strip(" .\t")
        amounts = DOLLAR_RE.findall(s)
        if not amounts:
            continue
        low = s.lower()
        is_client = any(w in low for w in CLIENT_CARD)
        is_mine = any(w in low for w in MY_CARD)
        card = "client" if is_client and not is_mine else "mine" if is_mine and not is_client else "unknown"
        if card == "unknown" and re.search(r"\bmiles?\b", low):
            notes.append(f"Ignored the dollar amount in a mileage sentence: \"{s}\"")
            continue
        for a in amounts:
            try:
                purchases.append(Purchase(day, s, Decimal(a.replace(",", "")), card))
            except InvalidOperation:
                pass
    return miles_found, purchases, notes


def parse_notes(text, year):
    """Return a list of Visit objects, one per 'Visit note' block."""
    blocks, current = [], None
    for line in text.splitlines():
        if NOTE_START.match(line):
            current = []
            blocks.append(current)
        elif current is not None:
            current.append(line)
    visits = []
    for i, lines in enumerate(blocks, 1):
        fields, key = defaultdict(list), None
        for line in lines:
            for k, rx in LABEL_RES:
                m = rx.match(line)
                if m:
                    key = k
                    fields[k].append(m.group(1))
                    break
            else:
                if key and line.strip():
                    fields[key].append(line.strip())
        v = Visit(label=f"note {i}")
        joined = {k: " ".join(x for x in vals if x).strip() for k, vals in fields.items()}

        day, weekday = parse_date(joined.get("date", ""), year) if joined.get("date") else (None, None)
        if day is None:
            v.errors.append("I could not read the date. Write it like: Tuesday, October 13")
        else:
            v.day = day
            v.label = f"note {i} ({pretty_date(day)})"
            if weekday and day.strftime("%A").lower() != weekday.lower():
                v.warnings.append(
                    f"{weekday} {day.strftime('%B')} {day.day} is not a {weekday} in {day.year}. "
                    "Check the date, or run with --year if it is another year.")

        times = parse_times(joined.get("times", "")) if joined.get("times") else None
        if times is None:
            v.errors.append("I could not read the start and end times. "
                            "Write them like: 10:00 am - 2:15 pm (with am or pm on both).")
        else:
            v.start, v.end = times
            if v.end <= v.start:
                v.errors.append("The end time is not after the start time. Check am and pm.")
            else:
                v.actual = v.end - v.start

        drives = joined.get("drives", "")
        if drives:
            miles_found, v.purchases, notes = parse_drives(drives, day)
            v.miles = sum(miles_found, Decimal("0"))
            v.warnings.extend(notes)
            if len(miles_found) > 1:
                parts = " + ".join(str(x) for x in miles_found)
                v.warnings.append(f"Found {len(miles_found)} mileage numbers ({parts}). "
                                  "Check none is counted twice.")
            if re.search(r"\bmiles?\b", drives, re.I) and not miles_found:
                v.warnings.append("The note mentions miles but I found no number.")
        visits.append(v)

    seen = {}
    for v in visits:
        if v.day and v.start is not None:
            k = (v.day, v.start, v.end)
            if k in seen:
                v.warnings.append(f"Looks like a repeat of {seen[k]}. Remove one if it is.")
            seen[k] = v.label
    return visits


# --------------------------------------------------------------------------
# The money
# --------------------------------------------------------------------------

@dataclass
class Line:
    day: date
    start: int
    end: int
    actual: int
    billed: int
    amount: Decimal
    miles: Decimal
    mileage_amount: Decimal


@dataclass
class Invoice:
    number: str
    client_key: str
    client_name: str
    payer: str
    sent: date
    due: date
    first: date
    last: date
    rate: Decimal
    min_hours: Decimal
    mileage_rate: Decimal
    lines: list
    purchases: list
    visits_total: Decimal
    mileage_total: Decimal
    repay_total: Decimal
    total: Decimal
    warnings: list
    unsure: list


def compute_invoice(visits, client, client_key, number, sent):
    rate = Decimal(str(client["rate"]))
    min_hours = Decimal(str(client["minimum_hours"]))
    mileage_rate = Decimal(str(client.get("mileage_rate", DEFAULT_MILEAGE_RATE)))
    min_minutes = int((min_hours * 60).to_integral_value(ROUND_HALF_UP))

    lines, purchases, warnings = [], [], []
    for v in sorted((x for x in visits if not x.errors), key=lambda x: (x.day, x.start)):
        billed = max(v.actual, min_minutes)
        amount = money(Decimal(billed) * rate / 60)
        mileage_amount = money(v.miles * mileage_rate)
        lines.append(Line(v.day, v.start, v.end, v.actual, billed, amount, v.miles, mileage_amount))
        purchases.extend(v.purchases)
        for w in v.warnings:
            warnings.append(f"{v.label}: {w}")
    if not lines:
        raise ValueError("No visits could be billed.")

    per_day = defaultdict(int)
    for ln in lines:
        per_day[ln.day] += ln.billed
    for d, mins in per_day.items():
        if mins > MAX_DAY_MINUTES:
            warnings.append(f"{pretty_date(d)}: {hm(mins)} billed in one day is over the "
                            "12 hours a day the agreement allows.")

    mine = [p for p in purchases if p.card == "mine"]
    for p in mine:
        if p.amount >= OWN_CARD_LIMIT:
            warnings.append(f"{pretty_date(p.date)}: {fmt(p.amount)} on my own card is not "
                            "under $100. The agreement says my card is a backup only under $100.")
    unsure = [p for p in purchases if p.card == "unknown"]

    visits_total = sum((ln.amount for ln in lines), Decimal("0"))
    mileage_total = sum((ln.mileage_amount for ln in lines), Decimal("0"))
    repay_total = sum((p.amount for p in mine), Decimal("0"))
    return Invoice(
        number=str(number), client_key=client_key, client_name=client["name"],
        payer=client.get("payer", ""), sent=sent, due=sent + timedelta(days=7),
        first=lines[0].day, last=lines[-1].day, rate=rate, min_hours=min_hours,
        mileage_rate=mileage_rate, lines=lines, purchases=purchases,
        visits_total=money(visits_total), mileage_total=money(mileage_total),
        repay_total=money(repay_total),
        total=money(visits_total + mileage_total + repay_total),
        warnings=warnings, unsure=unsure)


# --------------------------------------------------------------------------
# Showing the result
# --------------------------------------------------------------------------

def render_understood(inv, visits):
    out = ["WHAT I UNDERSTOOD (check this before saying yes)", ""]
    for ln in inv.lines:
        short = "  (shortest visit applies)" if ln.billed > ln.actual else ""
        out.append(f"{pretty_date(ln.day)}  {clock(ln.start)} - {clock(ln.end)}  "
                   f"worked {hm(ln.actual)}  billed {hm(ln.billed)}  {fmt(ln.amount)}{short}")
        if ln.miles:
            out.append(f"    miles {ln.miles} x {fmt(inv.mileage_rate)} = {fmt(ln.mileage_amount)}")
    if inv.purchases:
        out.append("")
        out.append("Purchases found:")
    for p in inv.purchases:
        where = {"client": "client's card, listed only, not added",
                 "mine": "my card, added to the total",
                 "unknown": "NOT ADDED: I could not tell whose card"}[p.card]
        out.append(f"    {pretty_date(p.date)}  {fmt(p.amount)} ({where}): {p.text}")
    out += ["", f"Visits {fmt(inv.visits_total)}   Mileage {fmt(inv.mileage_total)}   "
                f"My card to be repaid {fmt(inv.repay_total)}",
            f"TOTAL DUE {fmt(inv.total)}   (due {pretty_date(inv.due)})"]
    problems = [(v.label, e) for v in visits for e in v.errors]
    if problems:
        out += ["", "PROBLEMS (these visits are NOT in the invoice; fix the note and run again):"]
        out += [f"  - {lab}: {e}" for lab, e in problems]
    if inv.unsure:
        out += ["", "NEEDS YOUR DECISION (not added to the total):"]
        out += [f"  - {fmt(p.amount)} on {pretty_date(p.date)}: \"{p.text}\". "
                "Say \"your card\" or \"my card\" in the note." for p in inv.unsure]
    if inv.warnings:
        out += ["", "CHECK THESE:"] + [f"  - {w}" for w in inv.warnings]
    return "\n".join(out)


def render_invoice(inv, config):
    sender = config.get("from", "[Your full name], [Last Name] Home Support")
    L = [f"INVOICE {inv.number} - {sender.split(',', 1)[-1].strip() if ',' in sender else sender}",
         f"From: {sender}"]
    if config.get("contact"):
        L.append(f"Contact: {config['contact']}")
    L += [f"Client: {inv.client_name}"]
    if inv.payer:
        L.append(f"Paid by: {inv.payer}")
    week = (pretty_date(inv.first) if inv.first == inv.last
            else f"{pretty_date(inv.first)} to {pretty_date(inv.last)}")
    L += [f"Date sent: {pretty_date(inv.sent)}",
          f"Week covered: {week}",
          f"Payment due: {pretty_date(inv.due)} (within 7 days)",
          f"Hourly rate: {fmt(inv.rate)} an hour",
          f"Shortest visit: {inv.min_hours.normalize():f} hours", "",
          "VISITS (actual time worked, to the nearest minute; the shortest visit applies)",
          f"{'Date':<14}{'Time':<22}{'Billed':<12}{'Amount':>10}"]
    for ln in inv.lines:
        L.append(f"{pretty_date(ln.day):<14}{clock(ln.start) + ' - ' + clock(ln.end):<22}"
                 f"{hm(ln.billed):<12}{fmt(ln.amount):>10}")
    L += [f"{'Visits total':<48}{fmt(inv.visits_total):>10}", ""]

    drives = [ln for ln in inv.lines if ln.miles]
    if drives:
        L += [f"MILEAGE (drives with or for the client, {fmt(inv.mileage_rate)} a mile)",
              f"{'Date':<14}{'Miles':<10}{'Amount':>10}"]
        for ln in drives:
            L.append(f"{pretty_date(ln.day):<14}{str(ln.miles):<10}{fmt(ln.mileage_amount):>10}")
        L += [f"{'Mileage total':<24}{fmt(inv.mileage_total):>10}", ""]

    listed = [p for p in inv.purchases if p.card in ("client", "mine")]
    if listed:
        L += ["PURCHASES (receipts attached; the client's card is listed for the record only)"]
        for p in listed:
            card = "client's card" if p.card == "client" else "my card"
            L.append(f"{pretty_date(p.date)}  {fmt(p.amount):>9}  {card:<14}{p.text}")
        L += [f"{'On my card, to be repaid':<24}{fmt(inv.repay_total):>10}", ""]

    L += ["TOTAL DUE",
          f"{'Visits':<34}{fmt(inv.visits_total):>10}",
          f"{'Mileage':<34}{fmt(inv.mileage_total):>10}",
          f"{'Purchases on my card':<34}{fmt(inv.repay_total):>10}",
          f"{'TOTAL DUE THIS WEEK':<34}{fmt(inv.total):>10}",
          f"Please pay by {pretty_date(inv.due)}, by check, cash or Zelle.",
          "I give a receipt for every cash payment and every reimbursement."]
    return "\n".join(L) + "\n"


# --------------------------------------------------------------------------
# Tracker spreadsheet
# --------------------------------------------------------------------------

TRACKER_FIRST_ROW, TRACKER_LAST_ROW = 10, 59   # row 9 of the tracker is the example


def tracker_next_number(path):
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    used = sum(1 for r in range(TRACKER_FIRST_ROW, TRACKER_LAST_ROW + 1) if ws.cell(r, 1).value)
    return used + 1


def add_to_tracker(path, inv):
    """Fill the next empty tracker row. Formulas in the grey columns are left alone."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    for r in range(TRACKER_FIRST_ROW, TRACKER_LAST_ROW + 1):
        if not ws.cell(r, 1).value:
            week = (f"{inv.first.strftime('%b')} {inv.first.day}-{inv.last.day}"
                    if (inv.first.month == inv.last.month) else
                    f"{inv.first.strftime('%b')} {inv.first.day} - {inv.last.strftime('%b')} {inv.last.day}")
            ws.cell(r, 1).value = inv.number
            ws.cell(r, 2).value = inv.client_name
            ws.cell(r, 3).value = week
            ws.cell(r, 4).value = inv.sent
            ws.cell(r, 6).value = float(inv.total)
            wb.save(path)
            return r
    raise RuntimeError("The tracker has no empty rows left.")


# --------------------------------------------------------------------------
# Command line
# --------------------------------------------------------------------------

def load_config(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Turn visit notes into a weekly invoice.")
    ap.add_argument("notes", help="text file with the pasted visit notes for one client")
    ap.add_argument("--client", required=True, help="client key from clients.json")
    ap.add_argument("--clients", default="clients.json", help="path to clients.json (private)")
    ap.add_argument("--year", type=int, default=date.today().year,
                    help="year for dates written without one (default: this year)")
    ap.add_argument("--sent", help="date the invoice is sent, YYYY-MM-DD (default: today)")
    ap.add_argument("--number", help="invoice number (default: next one in the tracker, or 1)")
    ap.add_argument("--tracker", help="path to invoice_tracker.xlsx to add a row to")
    ap.add_argument("--outdir", default="invoices", help="folder for the invoice (default: invoices)")
    ap.add_argument("--yes", action="store_true", help="skip the question and write the files")
    args = ap.parse_args(argv)

    config = load_config(args.clients)
    clients = config.get("clients", {})
    if args.client not in clients:
        sys.exit(f"No client '{args.client}' in {args.clients}. Known: {', '.join(clients) or 'none'}")
    client = clients[args.client]

    text = Path(args.notes).read_text(encoding="utf-8", errors="replace")
    visits = parse_notes(text, args.year)
    if not visits:
        sys.exit("I found no 'Visit note' in that file.")

    sent = datetime.strptime(args.sent, "%Y-%m-%d").date() if args.sent else date.today()
    number = args.number
    if number is None:
        number = tracker_next_number(args.tracker) if args.tracker else 1
    try:
        inv = compute_invoice(visits, client, args.client, number, sent)
    except ValueError as e:
        print(render_understood_errors(visits))
        sys.exit(str(e))

    print(render_understood(inv, visits))
    blocked = any(v.errors for v in visits)
    if blocked:
        print("\nNothing was written, because some visits could not be read.")
        return 1
    if not args.yes:
        if input("\nDoes this look right? Write the invoice? [y/N] ").strip().lower() != "y":
            print("Nothing was written.")
            return 0

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    out = outdir / f"invoice_{args.client}_{inv.number}_{sent.isoformat()}.txt"
    if out.exists():
        sys.exit(f"{out} already exists. Nothing was written. Use --number to pick another.")
    out.write_text(render_invoice(inv, config), encoding="utf-8")
    print(f"\nWrote {out}")
    if args.tracker:
        row = add_to_tracker(args.tracker, inv)
        print(f"Added invoice {inv.number} to row {row} of {args.tracker}")
    print("Reminder: attach the receipts, and keep this folder out of GitHub.")
    return 0


def render_understood_errors(visits):
    return "\n".join(f"{v.label}: {e}" for v in visits for e in v.errors)


if __name__ == "__main__":
    sys.exit(main())
